"""Add a 0_Part_Recommendations sheet (per-part z-score recommendations) to
Scrap_Recommendations_Report_v2.xlsx, as the FIRST sheet.

Format:  one row per scrap part with columns
  Part ID | Part Description | Built Time | Tool Number | Scrap Reason | Recommendation

Recommendation = top-10 z-scored deviations (|z|>=1) vs same-tool GOOD baseline:
  [SEVERITY] <param> running HIGH/LOW (z) -> Increase/Reduce
"""
import os, sys, io, shutil
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
import numpy as np, pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = r"C:\Users\12345\Desktop\IAC\Vacforming"
XLSX = os.path.join(ROOT, "Scrap_Recommendations_Report_v2.xlsx")

# ---- Re-derive cleaned df (same recipe as final notebook) ------------------
p_raw = pd.read_csv(os.path.join(ROOT, "data", "VF_export.csv")).drop(columns=lambda c: c.startswith('Unnamed'), errors='ignore')
s_raw = pd.read_csv(os.path.join(ROOT, "data", "VF_Scrap_export.csv")).drop(columns=lambda c: c.startswith('Unnamed'), errors='ignore')
p_raw['BuiltTime'] = pd.to_datetime(p_raw['BuiltTime'])
s_raw['ScrapTime'] = pd.to_datetime(s_raw['ScrapTime'])

def parse_kv(s):
    out = {}
    if pd.isna(s): return out
    for it in str(s).split(','):
        if ':' in it:
            k, v = it.rsplit(':', 1)
            try: out[k.strip()] = float(v.strip())
            except: pass
    return out
ds1 = p_raw['DS1'].apply(parse_kv).apply(pd.Series)
ds2 = p_raw['DS2'].apply(parse_kv).apply(pd.Series)
ds3 = p_raw['DS3'].apply(parse_kv).apply(pd.Series)
df = pd.concat([p_raw[['ID','PartDesc','BuiltTime']], ds1, ds2, ds3], axis=1)
df = df.merge(s_raw[['ID','Desc']].rename(columns={'Desc':'defect'}), on='ID', how='left')
df['is_scrap'] = df['defect'].notna()

# Treat sentinel as missing (so it doesn't dominate z-scores)
df.loc[df['Machine Cycle Time'] == 9999, 'Machine Cycle Time'] = np.nan
# Drop columns that are constant -> z-score undefined
for c in ['Pick Material Pos','TTF Circuit 2 Temp']:
    if c in df.columns: df = df.drop(columns=[c])

PARAM_COLS = [c for c in df.columns if c not in
              ('ID','PartDesc','BuiltTime','defect','is_scrap','Tool Number')]

# ---- Compute per-tool GOOD baseline (mean, std) ----------------------------
good = df[~df['is_scrap']]
GOOD_STATS = {}
for tool in sorted(df['Tool Number'].dropna().unique()):
    sub = good[good['Tool Number']==tool]
    GOOD_STATS[int(tool)] = {p: (sub[p].mean(), sub[p].std()) for p in PARAM_COLS}

# Human-friendly names (operator-readable)
NICE = {
    'Temp. THT at Trigger': 'Top Heater Temp (THT)',
    'Temp. BHT at Trigger': 'Bottom Heater Temp (BHT)',
    'heating time (tens of seconds)': 'Heating Time',
    'Stetch Length Pick Up': 'Stretch Length Pick Up',
    'Stetch Length Stretch 1': 'Stretch Length Stretch 1',
    'Stetch Length Stretch 2': 'Stretch Length Stretch 2',
    'Stetch Length Stretch 3': 'Stretch Length Stretch 3',
    'Stetch Length Form': 'Stretch Length Form',
    'Vacuum Time Tool Cavity A': 'Vacuum Time Cavity A',
    'Vacuum Time Tool Cavity B': 'Vacuum Time Cavity B',
    'Graining Vacuum TTF Cavity 1': 'Graining Vacuum TTF Cavity 1',
    'Graining Vacuum TTF Cavity 2': 'Graining Vacuum TTF Cavity 2',
}

# ---- Recommendation builder per row ----------------------------------------
def recommend_row(row, top_k=10, z_min=1.0):
    tool = row.get('Tool Number')
    if pd.isna(tool) or int(tool) not in GOOD_STATS: return ''
    stats_t = GOOD_STATS[int(tool)]
    entries = []
    for p, (mu, sigma) in stats_t.items():
        if sigma is None or pd.isna(sigma) or sigma == 0: continue
        v = row.get(p)
        if pd.isna(v): continue
        z = (v - mu) / sigma
        if abs(z) < z_min: continue
        entries.append((z, p))
    entries.sort(key=lambda x: -abs(x[0]))
    out = []
    for z, p in entries[:top_k]:
        sev = 'CRITICAL' if abs(z) > 3 else ('HIGH' if abs(z) > 2 else 'MODERATE')
        direction = 'HIGH' if z > 0 else 'LOW'
        action = 'Reduce' if z > 0 else 'Increase'
        nice = NICE.get(p, p)
        out.append(f'[{sev}] {nice} running {direction} ({z:+.1f}) -> {action}')
    return '; '.join(out)

# ---- Build sheet -----------------------------------------------------------
scrap = df[df['is_scrap']].sort_values('BuiltTime').reset_index(drop=True)
sheet = pd.DataFrame({
    'Part ID':          scrap['ID'].astype('Int64'),
    'Part Description': scrap['PartDesc'],
    'Built Time':       scrap['BuiltTime'].dt.strftime('%Y-%m-%d %H:%M:%S'),
    'Tool Number':      scrap['Tool Number'].astype('Int64'),
    'Scrap Reason':     scrap['defect'],
    'Recommendation':   scrap.apply(recommend_row, axis=1),
})
print(f'Rows in new sheet: {len(sheet)}')
print('\nFirst 3 rows:')
for _, r in sheet.head(3).iterrows():
    print(f'\n  Part {r["Part ID"]} | {r["Part Description"]} | Tool {r["Tool Number"]} | {r["Scrap Reason"]}')
    print(f'  {r["Recommendation"]}')

# ---- Insert as FIRST sheet in the Excel ------------------------------------
# Strategy: load existing wb, create new sheet at index 0, write data, then save.
wb = load_workbook(XLSX)

NEW_NAME = '0_Part_Recommendations'
if NEW_NAME in wb.sheetnames:
    del wb[NEW_NAME]
ws = wb.create_sheet(NEW_NAME, 0)

headers = list(sheet.columns)
ws.append(headers)
for _, row in sheet.iterrows():
    ws.append([row[c] if not pd.isna(row[c]) else '' for c in headers])

# Style header
header_font = Font(bold=True, color='FFFFFF', size=11)
header_fill = PatternFill('solid', fgColor='2C3E50')
for c, _ in enumerate(headers, 1):
    cell = ws.cell(row=1, column=c)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal='center', vertical='center')

# Column widths
widths = {'A':10, 'B':28, 'C':22, 'D':8, 'E':32, 'F':140}
for col, w in widths.items(): ws.column_dimensions[col].width = w

# Wrap recommendation column
for r in range(2, ws.max_row + 1):
    ws.cell(row=r, column=6).alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[r].height = 60

# Freeze top row
ws.freeze_panes = 'A2'

wb.save(XLSX)
print(f'\nWrote sheet "{NEW_NAME}" as first tab in {XLSX}')
print(f'Sheets now: {wb.sheetnames}')
